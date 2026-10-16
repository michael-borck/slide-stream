"""Build styled.pptx: a deliberately distinctive design for --preserve-style.

The deck uses a dark navy background, an orange accent bar, coloured text,
and drawn shapes. If the video frames show plain white slides with black
text, the design was lost; if they look like this deck, it survived.
Every slide has speaker notes, which drive the narration.

    uv run python e2e/presenton/make_styled_deck.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x14, 0x21, 0x3D)
ORANGE = RGBColor(0xFC, 0xA3, 0x11)
CREAM = RGBColor(0xF4, 0xEF, 0xE6)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)

OUT = Path(__file__).with_name("styled.pptx")


def _background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = NAVY


def _accent_bar(slide, prs) -> None:
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.35), prs.slide_height
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ORANGE
    bar.line.fill.background()


def _text(slide, text, left, top, width, height, size, color, bold=False,
          align=PP_ALIGN.LEFT) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.alignment = align
        run = para.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "Georgia"
        para.space_after = Pt(size * 0.5)


def _notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def _slide(prs, title, bullets, notes):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _background(slide)
    _accent_bar(slide, prs)
    _text(slide, title, Inches(0.9), Inches(0.6), Inches(11.5), Inches(1.2),
          40, ORANGE, bold=True)
    if bullets:
        _text(slide, [f"\u2022  {b}" for b in bullets], Inches(0.9),
              Inches(2.0), Inches(7.2), Inches(4.8), 24, CREAM)
    _notes(slide, notes)
    return slide


def main() -> None:
    prs = Presentation()
    prs.slide_width = Emu(Inches(13.333))
    prs.slide_height = Emu(Inches(7.5))

    # 1. Title slide: centred, with a teal band.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _background(slide)
    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(4.6), prs.slide_width, Inches(0.25)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = TEAL
    band.line.fill.background()
    _text(slide, "Running a Good Tutorial", Inches(0.8), Inches(2.4),
          Inches(11.7), Inches(1.4), 54, ORANGE, bold=True, align=PP_ALIGN.CENTER)
    _text(slide, "Five habits for first-time tutors", Inches(0.8), Inches(5.1),
          Inches(11.7), Inches(0.8), 26, CREAM, align=PP_ALIGN.CENTER)
    _notes(slide, "Welcome. This short session is for anyone running their "
           "first tutorial this semester. We'll cover five habits that make "
           "the hour work for students, and for you.")

    _slide(prs, "Arrive With One Goal",
           ["Pick the single thing students should leave able to do",
            "Write it on the board before they arrive",
            "Check against it at the end"],
           "Before anything else, decide what one thing students should be "
           "able to do when they walk out. Put it on the board where everyone "
           "can see it, and come back to it in the last five minutes.")

    _slide(prs, "Talk Less Than Half the Time",
           ["Explain for ten minutes at most",
            "Then give them a problem to work on",
            "Walk the room and listen"],
           "A tutorial is not a second lecture. Keep your explanation short, "
           "then hand over a problem. Most of the learning happens while they "
           "work and you move around the room listening.")

    # 4. A diagram slide made of shapes: a three-step cycle.
    slide = _slide(prs, "Ask, Wait, Follow Up", [],
                   "When you ask a question, wait. Count to five in your head "
                   "before you say anything. Then follow up on the answer, "
                   "even a wrong one, and ask how they got there.")
    for i, (label, colour) in enumerate(
        [("Ask", TEAL), ("Wait", ORANGE), ("Follow up", TEAL)]
    ):
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(1.3 + i * 4.0), Inches(2.6),
            Inches(2.6), Inches(2.6),
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = colour
        circle.line.fill.background()
        frame = circle.text_frame
        frame.text = label
        run = frame.paragraphs[0].runs[0]
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = NAVY
        frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        if i < 2:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW, Inches(4.05 + i * 4.0), Inches(3.6),
                Inches(1.1), Inches(0.6),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = CREAM
            arrow.line.fill.background()

    _slide(prs, "Make Mistakes Normal",
           ["Share a mistake you made learning this",
            "Thank people for wrong answers",
            "Fix errors together, out loud"],
           "Students will only try things if getting it wrong feels safe. "
           "Tell them about a mistake you made with this material. When "
           "someone gets an answer wrong, thank them, and work through it "
           "together.")

    _slide(prs, "Close the Loop",
           ["Return to the goal on the board",
            "One minute: each student writes what they learned",
            "Read the cards before next week"],
           "Finish where you started. Point back at the goal, then give "
           "everyone one minute to write down what they learned. Read those "
           "cards before next week. They'll tell you what to repeat.")

    prs.save(str(OUT))
    print(f"Wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
