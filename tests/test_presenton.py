"""Tests for the optional Presenton integration (deck design) and the
designed-slides-to-video path (PDF rasterisation).

Presenton (https://github.com/presenton/presenton, Apache-2.0) is an external
service: every HTTP call here is faked, so the suite stays hermetic.
"""

# python-pptx ships incomplete type stubs (same note as test_cli.py).
# pyright: reportArgumentType=false, reportAttributeAccessIssue=false, reportOptionalMemberAccess=false

import json
import sys
import wave
from pathlib import Path

import pytest
import requests
from PIL import Image
from pptx import Presentation
from typer.testing import CliRunner

from slide_stream.cli import app
from slide_stream.export_pptx import export_pptx
from slide_stream.presenton import (
    DEFAULT_PRESENTON_INSTRUCTIONS,
    PresentonClient,
    PresentonError,
    inject_speaker_notes,
    resolve_presenton_settings,
    slides_to_markdown,
)
from slide_stream.slide_render import (
    SlideRenderError,
    pdf_to_slide_images,
    pptx_to_pdf,
)

SLIDES = [
    {"title": "Intro", "content": ["Welcome", "Agenda"], "has_real_title": True},
    {"title": "Slide 2", "content": ["Untitled point"], "has_real_title": False},
    {"title": "", "content": [], "has_real_title": False},
]

DECK_MD = "# Alpha\n\n- One\n- Two\n\n# Beta\n\n- Three\n"

FAST_VIDEO_CONFIG = (
    "settings:\n"
    "  video:\n"
    "    resolution: [320, 180]\n"
    "    fps: 8\n"
    "    default_slide_duration: 1.0\n"
    "    slide_duration_padding: 0.2\n"
)


def _write_silent_wav(filename, seconds=0.3):
    with wave.open(str(filename), "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(22050)
        w.writeframes(b"\x00\x00" * int(22050 * seconds))


# --- slides_to_markdown / settings ---------------------------------------


def test_slides_to_markdown_keeps_slide_count_and_skips_placeholder_titles():
    out = slides_to_markdown(SLIDES)
    assert len(out) == 3  # 1:1 with the source slides: narration alignment
    assert out[0] == "# Intro\n\n- Welcome\n- Agenda"
    assert out[1] == "- Untitled point"  # parser's "Slide 2" is not a real title
    assert out[2] == "# Slide 3"  # empty slide still gets an entry


def test_slides_to_markdown_sends_markdown_decks_verbatim():
    """Code blocks, tables and sub-headings reach Presenton as Markdown,
    not flattened into bullets."""
    from slide_stream.parser import parse_markdown

    md = (
        "# Code\n\n```python\nx = 1\n```\n\n"
        "# Table\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n"
        "# Parts\n\n## Part A\n\n- point\n\n"
        "# Empty\n"
    )
    out = slides_to_markdown(parse_markdown(md))
    assert out == [
        "# Code\n\n```python\nx = 1\n```",
        "# Table\n\n| a | b |\n|---|---|\n| 1 | 2 |",
        "# Parts\n\n## Part A\n\n- point",
        "# Empty",
    ]


def test_resolve_settings_cli_wins_over_config_and_env(monkeypatch):
    monkeypatch.setenv("PRESENTON_API_KEY", "env-key")
    config = {
        "providers": {
            "presenton": {
                "base_url": "http://config:5001",
                "api_key": "config-key",
                "template": "modern",
            }
        }
    }
    settings = resolve_presenton_settings(config, url="http://cli:9000/")
    assert settings["base_url"] == "http://cli:9000"  # trailing slash trimmed
    assert settings["api_key"] == "config-key"  # config beats env
    assert settings["template"] == "modern"

    settings = resolve_presenton_settings({}, template="swift")
    assert settings["base_url"] == "http://localhost:5001"
    assert settings["api_key"] == "env-key"  # env is the last resort
    assert settings["template"] == "swift"


# --- HTTP client (faked session) ------------------------------------------


class _FakeResponse:
    def __init__(self, status_code=200, payload=None, content=b""):
        self.status_code = status_code
        self._payload = payload or {}
        self.content = content
        self.text = ""

    def json(self):
        return self._payload

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.HTTPError(f"{self.status_code}", response=self)


class _FakeSession:
    """Records calls; serves queued responses in order."""

    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []

    def post(self, url, json=None, headers=None, timeout=None):
        self.calls.append(("POST", url, json, headers))
        return self.responses.pop(0)

    def get(self, url, headers=None, timeout=None):
        self.calls.append(("GET", url, None, headers))
        return self.responses.pop(0)


def test_generate_deck_posts_slides_markdown_and_downloads(tmp_path):
    session = _FakeSession([
        _FakeResponse(payload={
            "presentation_id": "abc-123",
            "path": "/app_data/abc-123/Deck.pptx",
            "edit_path": "/presentation?id=abc-123",
        }),
        _FakeResponse(content=b"PPTX-BYTES"),
    ])
    client = PresentonClient(
        "http://presenton:5001", api_key="sk-test", session=session
    )
    dest = tmp_path / "out.pptx"
    result = client.generate_deck(SLIDES[:1], "general", "pptx", dest)

    assert dest.read_bytes() == b"PPTX-BYTES"
    assert result.presentation_id == "abc-123"
    method, url, payload, headers = session.calls[0]
    assert (method, url) == (
        "POST", "http://presenton:5001/api/v1/ppt/presentation/generate"
    )
    assert payload["content"] == ""
    assert payload["slides_markdown"] == ["# Intro\n\n- Welcome\n- Agenda"]
    assert payload["template"] == "general"
    assert payload["export_as"] == "pptx"
    # No added title/TOC slides: the deck must stay 1:1 with the source.
    assert payload["include_title_slide"] is False
    assert payload["include_table_of_contents"] is False
    assert headers["Authorization"] == "Bearer sk-test"
    # Server-relative path is downloaded from the Presenton origin.
    assert session.calls[1][1] == "http://presenton:5001/app_data/abc-123/Deck.pptx"
    # The faithfulness instructions go with every request by default.
    assert payload["instructions"] == DEFAULT_PRESENTON_INSTRUCTIONS


def test_default_instructions_cover_the_trial_failures():
    """The live trial saw code replaced by prose and invented people/rows."""
    text = DEFAULT_PRESENTON_INSTRUCTIONS.lower()
    assert "code" in text and "exactly" in text
    assert "do not invent" in text


def test_empty_instructions_are_not_sent(tmp_path):
    session = _FakeSession([
        _FakeResponse(payload={"presentation_id": "p", "path": "/d/x.pptx"}),
        _FakeResponse(content=b"X"),
    ])
    client = PresentonClient("http://p", session=session, instructions="")
    client.generate_deck(SLIDES[:1], "general", "pptx", tmp_path / "x.pptx")
    assert "instructions" not in session.calls[0][2]


def test_resolve_instructions_flag_config_default():
    assert (
        resolve_presenton_settings({})["instructions"]
        == DEFAULT_PRESENTON_INSTRUCTIONS
    )
    config = {"providers": {"presenton": {"instructions": " Keep it short. "}}}
    assert resolve_presenton_settings(config)["instructions"] == "Keep it short."
    assert resolve_presenton_settings(config, instructions="Flag wins.")[
        "instructions"] == "Flag wins."
    # An empty string is a deliberate "send nothing", not "use the default".
    off = {"providers": {"presenton": {"instructions": ""}}}
    assert resolve_presenton_settings(off)["instructions"] == ""


def test_generate_deck_pair_exports_pdf_for_same_presentation(tmp_path):
    session = _FakeSession([
        _FakeResponse(payload={"presentation_id": "p1", "path": "/d/deck.pptx"}),
        _FakeResponse(content=b"PPTX"),
        _FakeResponse(payload={"presentation_id": "p1", "path": "/d/deck.pdf"}),
        _FakeResponse(content=b"PDF"),
    ])
    client = PresentonClient("http://presenton:5001", session=session)
    pptx_dest, pdf_dest = tmp_path / "deck.pptx", tmp_path / "deck.pdf"
    client.generate_deck_pair(SLIDES[:1], "general", pptx_dest, pdf_dest)

    assert pptx_dest.read_bytes() == b"PPTX"
    assert pdf_dest.read_bytes() == b"PDF"
    export_calls = [c for c in session.calls if c[1].endswith("/p1/export")]
    assert len(export_calls) == 1
    assert export_calls[0][2] == {"export_as": "pdf"}


def test_generate_deck_http_error_surfaces_server_detail(tmp_path):
    session = _FakeSession([
        _FakeResponse(status_code=400, payload={"detail": "Template not found."})
    ])
    client = PresentonClient("http://presenton:5001", session=session)
    with pytest.raises(PresentonError, match="Template not found"):
        client.generate_deck(SLIDES[:1], "nope", "pptx", tmp_path / "x.pptx")


def test_generate_deck_unreachable_server_explains_how_to_start(tmp_path):
    class DeadSession:
        def post(self, url, json=None, headers=None, timeout=None):
            raise requests.ConnectionError("refused")

    client = PresentonClient("http://localhost:5999", session=DeadSession())
    with pytest.raises(PresentonError, match="could not reach Presenton"):
        client.generate_deck(SLIDES[:1], "general", "pptx", tmp_path / "x.pptx")


# --- speaker notes ---------------------------------------------------------


def test_inject_speaker_notes_writes_in_slide_order(tmp_path):
    deck = tmp_path / "deck.pptx"
    export_pptx(
        [
            {"title": "A", "content": ["one"], "notes": "Old A"},
            {"title": "B", "content": ["two"], "notes": "Old B"},
        ],
        deck,
    )
    # A blank entry leaves that slide's existing notes untouched.
    inject_speaker_notes(deck, ["Speak A now.", ""])
    prs = Presentation(str(deck))
    notes = [s.notes_slide.notes_text_frame.text.strip() for s in prs.slides]
    assert notes == ["Speak A now.", "Old B"]


# --- PDF rasterisation -----------------------------------------------------


def _blank_pdf(path, pages, width=1280, height=720):
    pdfium = pytest.importorskip("pypdfium2")
    doc = pdfium.PdfDocument.new()
    for _ in range(pages):
        doc.new_page(width, height)
    doc.save(str(path))
    doc.close()
    return path


def test_pdf_to_slide_images_without_pypdfium2(tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "pypdfium2", None)
    with pytest.raises(SlideRenderError, match="slide-stream\\[presenton\\]"):
        pdf_to_slide_images(tmp_path / "x.pdf", tmp_path / "out")


def test_pdf_to_slide_images_renders_every_page(tmp_path):
    pdf = _blank_pdf(tmp_path / "slides.pdf", pages=2)

    images = pdf_to_slide_images(pdf, tmp_path / "imgs", width=640)
    assert [p.name for p in images] == ["slide_1.png", "slide_2.png"]
    with Image.open(images[0]) as im:
        assert im.size == (640, 360)
        assert im.mode == "RGB"


def test_pdf_to_slide_images_exact_width_despite_rounding(tmp_path):
    # 1000pt page at width 1919 gives a fractional scale; output must still
    # be exactly 1919 px wide for the video frame.
    pdf = _blank_pdf(tmp_path / "odd.pdf", pages=1, width=1000, height=563)
    (image,) = pdf_to_slide_images(pdf, tmp_path / "imgs", width=1919)
    with Image.open(image) as im:
        assert im.width == 1919


def test_pdf_to_slide_images_unreadable_pdf(tmp_path):
    pytest.importorskip("pypdfium2")
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"not a pdf")
    with pytest.raises(SlideRenderError, match="Could not open"):
        pdf_to_slide_images(bad, tmp_path / "imgs")


def test_pptx_to_pdf_without_libreoffice(tmp_path, monkeypatch):
    monkeypatch.setattr("slide_stream.slide_render.shutil.which", lambda _: None)
    with pytest.raises(SlideRenderError, match="LibreOffice"):
        pptx_to_pdf(tmp_path / "deck.pptx", tmp_path)


# --- CLI: `presenton` command ----------------------------------------------


class _FakeClient:
    """Stands in for PresentonClient: copies a prepared deck to the dest."""

    last = None

    def __init__(self, source: Path):
        self.source = source
        self.calls = []
        _FakeClient.last = self

    def generate_deck(self, slides, template, export_as, dest):
        self.calls.append((template, export_as, len(slides)))
        Path(dest).write_bytes(self.source.read_bytes())
        return type("R", (), {"presentation_id": "p1", "edit_url": None})()


def test_cli_presenton_command_writes_deck_with_notes(tmp_path, monkeypatch):
    template_deck = tmp_path / "designed.pptx"
    export_pptx(
        [{"title": "Alpha", "content": []}, {"title": "Beta", "content": []}],
        template_deck,
    )
    monkeypatch.setattr(
        "slide_stream.cli._presenton_client",
        lambda settings: _FakeClient(template_deck),
    )
    md = tmp_path / "deck.md"
    md.write_text(DECK_MD)
    out = tmp_path / "out.pptx"

    result = CliRunner().invoke(app, ["presenton", str(md), str(out)])
    assert result.exit_code == 0, result.output
    assert _FakeClient.last.calls == [("general", "pptx", 2)]
    prs = Presentation(str(out))
    first_notes = prs.slides[0].notes_slide.notes_text_frame.text
    assert "Alpha" in first_notes and "One" in first_notes

    # Refuses to clobber an existing deck without --force.
    result = CliRunner().invoke(app, ["presenton", str(md), str(out)])
    assert result.exit_code != 0


# --- CLI: `create` with designed slides ------------------------------------


def _pngs(tmp_path: Path, count: int) -> list[Path]:
    paths = []
    for i in range(count):
        p = tmp_path / f"full_{i + 1}.png"
        Image.new("RGB", (320, 180), (10 * (i + 1), 0, 0)).save(p)
        paths.append(p)
    return paths


def test_cli_create_slides_pdf_uses_full_slide_images(
    tmp_path, mocker, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    md = tmp_path / "deck.md"
    md.write_text(DECK_MD)
    config_file = tmp_path / "config.yaml"
    config_file.write_text(FAST_VIDEO_CONFIG)

    full = _pngs(tmp_path, 2)
    monkeypatch.setattr(
        "slide_stream.cli.pdf_to_slide_images", lambda *a, **k: full
    )
    # The artwork provider must not be asked for images in full-slide mode.
    provider = mocker.MagicMock()
    provider.name = "swarmui"
    provider.generate_image.side_effect = AssertionError("provider used")
    mocker.patch(
        "slide_stream.cli.ProviderFactory.create_image_provider",
        return_value=provider,
    )
    fake_tts = mocker.MagicMock()
    fake_tts.save.side_effect = _write_silent_wav
    mocker.patch("gtts.gTTS", return_value=fake_tts)

    output = tmp_path / "out.mp4"
    result = CliRunner().invoke(app, [
        "create", str(md), str(output),
        "--config", str(config_file),
        "--slides-pdf", str(tmp_path / "slides.pdf"),
    ])
    assert result.exit_code == 0, result.output
    assert output.exists()


def test_cli_create_slides_pdf_count_mismatch_fails(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    md = tmp_path / "deck.md"
    md.write_text(DECK_MD)
    monkeypatch.setattr(
        "slide_stream.cli.pdf_to_slide_images",
        lambda *a, **k: _pngs(tmp_path, 1),  # 1 page for a 2-slide deck
    )
    result = CliRunner().invoke(app, [
        "create", str(md), str(tmp_path / "out.mp4"),
        "--slides-pdf", str(tmp_path / "slides.pdf"),
    ])
    assert result.exit_code != 0
    assert "does not match" in result.output


def test_cli_create_template_rejected_for_pptx_output(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    md = tmp_path / "deck.md"
    md.write_text(DECK_MD)
    result = CliRunner().invoke(app, [
        "create", str(md), str(tmp_path / "out.pptx"),
        "--template", "general",
    ])
    assert result.exit_code != 0
    assert "presenton" in result.output.lower()


# --- ADR-0013: fidelity report and rebuilt code slides ---------------------

CODE_DECK_MD = (
    "# Intro\n\n- welcome\n\n"
    "# Code\n\n```python\ndef f(x):\n    return x + 1\n```\n\n- adds one\n"
)


class _DesignClient:
    """A fake Presenton that designs a deck which paraphrases the code."""

    def __init__(self, tmp_path: Path):
        self.tmp = tmp_path
        self.exports: list[str] = []

    def _deck(self, dest):
        export_pptx(
            [{"title": "Intro", "content": ["welcome"]},
             {"title": "Code", "content": ["A function that adds one."]}],
            Path(dest),
        )

    def _pdf(self, dest, texts=("Intro welcome", "Code A function that adds one.")):
        pdfium = pytest.importorskip("pypdfium2")
        doc = pdfium.PdfDocument.new()
        for _ in texts:
            doc.new_page(320, 180)
        doc.save(str(dest))
        doc.close()

    def generate_deck(self, slides, template, export_as, dest):
        if export_as == "pptx":
            self._deck(dest)
        else:
            self._pdf(dest)
        return type("R", (), {"presentation_id": "p1", "edit_url": None})()

    def export_deck(self, presentation_id, export_as, dest):
        self.exports.append(export_as)
        self._pdf(dest)

    def generate_deck_pair(self, slides, template, pptx_dest, pdf_dest):
        self._deck(pptx_dest)
        self._pdf(pdf_dest)
        return type("R", (), {"presentation_id": "p1", "edit_url": None})()


def _slide_texts(path):
    from slide_stream.fidelity import pptx_slide_texts

    return pptx_slide_texts(path)


def test_cli_presenton_rebuilds_lost_code_and_writes_report(tmp_path, monkeypatch):
    pytest.importorskip("pypdfium2")
    client = _DesignClient(tmp_path)
    monkeypatch.setattr("slide_stream.cli._presenton_client", lambda settings: client)
    # The blank fake PDF has no text, so read the .pptx instead.
    monkeypatch.setattr("slide_stream.presenton.designed_texts",
                        lambda pptx, pdf: _slide_texts(pptx))
    md = tmp_path / "deck.md"
    md.write_text(CODE_DECK_MD)
    out = tmp_path / "out.pptx"

    result = CliRunner().invoke(app, ["presenton", str(md), str(out)])
    assert result.exit_code == 0, result.output
    assert "slide rebuilt by SlideStream" in result.output
    assert client.exports == ["pdf"]  # for colours and text

    texts = _slide_texts(out)
    assert "def f(x):\n    return x + 1" in texts[1]
    assert "A function that adds one." not in texts[1]
    assert texts[0].startswith("Intro")  # untouched
    notes = [s.notes_slide.notes_text_frame.text for s in Presentation(str(out)).slides]
    assert "Code" in notes[1]  # narration still written to the notes

    report = json.loads((tmp_path / "out.fidelity.json").read_text())
    assert report["code_slides_rebuilt"] == [2]
    assert report["slides"][1]["code_kept"] is False


def test_cli_presenton_strict_fails_when_code_cannot_be_fixed(tmp_path, monkeypatch):
    client = _DesignClient(tmp_path)
    monkeypatch.setattr("slide_stream.cli._presenton_client", lambda settings: client)
    monkeypatch.setattr("slide_stream.presenton.designed_texts",
                        lambda pptx, pdf: _slide_texts(pptx))
    md = tmp_path / "deck.md"
    md.write_text(CODE_DECK_MD)
    cfg = tmp_path / "c.yaml"
    cfg.write_text("providers:\n  presenton:\n    fidelity: strict\n    code_slides: never\n")

    result = CliRunner().invoke(app, [
        "presenton", str(md), str(tmp_path / "out.pptx"), "--config", str(cfg)])
    assert result.exit_code == 1
    assert "Strict fidelity check failed" in result.output
    assert "code missing on slide(s) 2" in result.output


def test_cli_presenton_rejects_bad_fidelity_setting(tmp_path, monkeypatch):
    monkeypatch.setattr("slide_stream.cli._presenton_client",
                        lambda settings: _DesignClient(tmp_path))
    md = tmp_path / "deck.md"
    md.write_text(CODE_DECK_MD)
    cfg = tmp_path / "c.yaml"
    cfg.write_text("providers:\n  presenton:\n    code_slides: sometimes\n")
    result = CliRunner().invoke(app, [
        "presenton", str(md), str(tmp_path / "out.pptx"), "--config", str(cfg)])
    assert result.exit_code == 1
    assert "code_slides must be one of auto, always, never" in result.output


def test_cli_create_template_swaps_code_slide_frame_and_deck(
    tmp_path, mocker, monkeypatch
):
    pytest.importorskip("pypdfium2")
    monkeypatch.chdir(tmp_path)
    md = tmp_path / "deck.md"
    md.write_text(CODE_DECK_MD)
    config_file = tmp_path / "config.yaml"
    config_file.write_text(FAST_VIDEO_CONFIG)
    client = _DesignClient(tmp_path)
    monkeypatch.setattr("slide_stream.cli._presenton_client", lambda settings: client)
    monkeypatch.setattr("slide_stream.presenton.designed_texts",
                        lambda pptx, pdf: _slide_texts(pptx))
    rendered = []
    from slide_stream import presenton as presenton_module

    real_render = presenton_module.render_code_slide

    def spy(dest, slide, theme, size):
        rendered.append((slide["title"], size))
        return real_render(dest, slide, theme, size)

    monkeypatch.setattr("slide_stream.presenton.render_code_slide", spy)
    fake_tts = mocker.MagicMock()
    fake_tts.save.side_effect = _write_silent_wav
    mocker.patch("gtts.gTTS", return_value=fake_tts)

    output = tmp_path / "talk.mp4"
    result = CliRunner().invoke(app, [
        "create", str(md), str(output), "--config", str(config_file),
        "--template", "general",
    ])
    assert result.exit_code == 0, result.output
    assert output.exists()
    assert rendered == [("Code", (320, 180))]  # one frame, at the frame size
    deck = tmp_path / "talk.pptx"
    assert "def f(x):" in _slide_texts(deck)[1]
    report = json.loads((tmp_path / "talk.fidelity.json").read_text())
    assert report["code_slides_rebuilt"] == [2]
