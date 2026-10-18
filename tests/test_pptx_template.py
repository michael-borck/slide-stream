"""Tests for designing decks from a user's PowerPoint template (ADR-0014)."""

import zipfile
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation

from slide_stream.code_render import highlight_lines
from slide_stream.pptx_template import (
    PANDOC_NAMES,
    build_deck,
    check_template,
    load_template,
    resolve_layouts,
)


def _texts(slide: Any) -> list[str]:
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]


def _shape_with(slide: Any, needle: str) -> Any:
    for sh in slide.shapes:
        tf: Any = getattr(sh, "text_frame", None)
        if tf is not None and needle in tf.text:
            return sh
    raise AssertionError(f"no shape containing {needle!r}")


def _table_of(slide: Any) -> Any:
    for sh in slide.shapes:
        tbl: Any = getattr(sh, "table", None)
        if tbl is not None:
            return tbl
    raise AssertionError("no table on the slide")


def _notes(slide: Any) -> str:
    return slide.notes_slide.notes_text_frame.text


def _template(tmp_path: Path) -> Path:
    tpl = tmp_path / "tpl.pptx"
    Presentation().save(str(tpl))
    return tpl


def test_default_template_has_the_standard_layout_names():
    """python-pptx's built-in deck carries pandoc's layout names."""
    names = {ly.name.casefold() for ly in Presentation().slide_layouts}
    for role, std in PANDOC_NAMES.items():
        assert std[0] in names, f"{role}: {std[0]!r} missing"


def test_resolve_layouts_matches_standard_names():
    layouts, missing = resolve_layouts(Presentation())
    assert missing == []
    assert layouts["title_slide"].name == "Title Slide"
    assert layouts["two_content"].name == "Two Content"


def test_renamed_layouts_are_matched_by_placeholder_signature(tmp_path):
    prs = Presentation()
    for ly in prs.slide_layouts:
        ly.name = f"Corporate {ly.name}"
    prs.save(str(tmp_path / "house.pptx"))
    renamed: Any = load_template(tmp_path / "house.pptx")[0]
    layouts, missing = resolve_layouts(renamed)
    assert set(layouts) == {"title_slide", "section_header", "title_content",
                            "two_content", "title_only", "blank"}
    # The names no longer match, so every standard name needed a stand-in.
    assert set(missing) == set(PANDOC_NAMES) - {"blank"}


def test_build_deck_fills_placeholders_and_notes(tmp_path):
    out = tmp_path / "deck.pptx"
    report = build_deck(_entries := [
        {"title": "Hash Tables", "content": ["Week 6, data structures"],
         "notes": "Welcome narration."},
        {"title": "The Problem", "content": ["One million records", "Half a million comparisons"],
         "notes": ""},
        {"title": "A Hash Function", "content": ["- hash() turns a key into an integer"],
         "code_blocks": [{"lang": "python",
                          "code": "def bucket_for(key: str, n: int) -> int:\n    return hash(key) % n"}],
         "notes": "Code narration."},
        {"title": "How Fast", "content": [],
         "tables": [[["Operation", "List", "Hash"], ["Look up", "O(n)", "O(1)"]]],
         "notes": ""},
    ], _template(tmp_path), out)
    prs = Presentation(str(out))
    assert len(prs.slides) == 4
    assert report.overflow == []

    first = prs.slides[0]
    assert "Hash Tables" in _texts(first)
    assert _notes(first) == "Welcome narration."
    assert first.slide_layout.name == "Title Slide"



def test_build_deck_code_slide_has_editable_coloured_code(tmp_path):
    out = tmp_path / "deck.pptx"
    build_deck([
        {"title": "A Hash Function", "content": [],
         "code_blocks": [{"lang": "python",
                          "code": "def bucket_for(key: str, n: int) -> int:\n    return hash(key) % n"}]},
    ], _template(tmp_path), out)
    slide = Presentation(str(out)).slides[0]
    code_box: Any = _shape_with(slide, "bucket_for")
    runs = [r for p in code_box.text_frame.paragraphs for r in p.runs if r.text.strip()]
    assert {r.font.name for r in runs} == {"Courier New"}
    assert len({str(r.font.color.rgb) for r in runs}) >= 3


def test_build_deck_table_is_a_native_table(tmp_path):
    out = tmp_path / "deck.pptx"
    build_deck([
        {"title": "How Fast", "content": [],
         "tables": [[["Operation", "List", "Hash"], ["Look up", "O(n)", "O(1)"]]]},
    ], _template(tmp_path), out)
    slide = Presentation(str(out)).slides[0]
    table: Any = _table_of(slide)
    assert table.cell(0, 0).text == "Operation"
    assert table.cell(1, 2).text == "O(1)"


def test_build_deck_shrinks_overflowing_text_and_reports_it(tmp_path):
    entries = [{"title": "Dense",
                "content": [f"Point {i} " + "word " * 30 for i in range(12)],
                "notes": ""}]
    out = tmp_path / "deck.pptx"
    report = build_deck(entries, _template(tmp_path), out)
    prs = Presentation(str(out))
    body: Any = _shape_with(prs.slides[0], "Point 0")
    sizes = {r.font.size.pt for p in body.text_frame.paragraphs
             for r in p.runs if r.font.size}
    assert sizes and max(sizes) < 18.0  # shrunk below the layout's size
    assert report.overflow and "slide 1" in report.overflow[0]


def test_potx_template_opens(tmp_path):
    """A .potx is a .pptx with a different content type; it must load."""
    src = _template(tmp_path)
    potx = tmp_path / "t.potx"
    with zipfile.ZipFile(src) as zin, \
            zipfile.ZipFile(potx, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in zin.namelist():
            data = zin.read(name)
            if name == "[Content_Types].xml":
                data = data.replace(
                    b"presentationml.presentation.main+xml",
                    b"presentationml.template.main+xml",
                )
            zout.writestr(name, data)
    _prs, kind = load_template(potx)
    assert kind == "potx"
    report = check_template(potx)
    assert report.template_kind == "potx"
    assert "title_slide" in report.layouts


def test_check_template_reports_renamed_layouts(tmp_path):
    prs = Presentation()
    for ly in prs.slide_layouts:
        ly.name = f"House {ly.name}"
    prs.save(str(tmp_path / "house.pptx"))
    report = check_template(tmp_path / "house.pptx")
    assert set(report.layouts) == {"title_slide", "section_header", "title_content",
                                   "two_content", "title_only", "blank"}
    assert set(report.missing) == set(PANDOC_NAMES) - {"blank"}


def test_missing_template_file(tmp_path):
    with pytest.raises(FileNotFoundError):
        build_deck([], tmp_path / "nope.pptx", tmp_path / "out.pptx")


def test_restyled_template_still_serves_every_role(tmp_path):
    """Colours on the layouts don't stop the roles from being served."""
    from pptx.dml.color import RGBColor

    prs = Presentation()
    for ly in prs.slide_layouts:
        for ph in ly.placeholders:
            if ph.has_text_frame:
                for para in ph.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.color.rgb = RGBColor(0x10, 0x30, 0x60)
    tpl = tmp_path / "styled.pptx"
    prs.save(str(tpl))
    out = tmp_path / "deck.pptx"
    report = build_deck(_entries := [
        {"title": "Hash Tables", "content": ["Week 6"], "notes": ""},
        {"title": "The Problem", "content": ["A", "B"], "notes": ""},
        {"title": "Code", "content": [],
         "code_blocks": [{"lang": "python", "code": "x = hash(1)"}], "notes": ""},
        {"title": "Table", "content": [],
         "tables": [[["a", "b"], ["1", "2"]]], "notes": ""},
    ], tpl, out)
    assert report.overflow == []
    assert len(Presentation(str(out)).slides) == 4


def test_highlight_language_reaches_the_code_box():
    lines = highlight_lines("x = 1", "python", dark=True)
    assert lines and lines[0][0].text.startswith("x")
