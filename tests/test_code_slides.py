"""Tests for SlideStream-built code slides in designed decks (ADR-0013)."""

from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

from slide_stream.code_slides import (
    choose_code_slides,
    neighbour_frames,
    plain_theme,
    render_code_slide,
    rewrite_pptx_code_slides,
    sample_theme,
    theme_from,
)
from slide_stream.config_loader import DEFAULT_CONFIG
from slide_stream.parser import parse_markdown

DECK = parse_markdown(
    "# Intro\n\n- welcome\n\n"
    "# Code\n\n```python\ndef f(x):\n    return x + 1\n```\n\n- adds one\n\n"
    "# More code\n\n```js\nconsole.log(1)\n```\n"
)
CODE = DECK[1]


def _frame(path: Path, bg, accent=None) -> Path:
    img = Image.new("RGB", (320, 180), bg)
    if accent:
        img.paste(Image.new("RGB", (60, 20), accent), (40, 60))
    img.save(path)
    return path


def test_sample_theme_takes_background_and_accent_from_frames(tmp_path):
    frames = [_frame(tmp_path / f"{i}.png", (248, 248, 248), (150, 60, 230)) for i in range(2)]
    theme = sample_theme(frames)
    assert theme is not None
    assert theme.bg == (248, 248, 248)  # exact, not rounded
    assert theme.accent == (144, 56, 224)  # quantised to steps of 8
    assert theme.fg == (34, 34, 34)  # light slide, dark text
    assert theme.dark is True  # dark code box on light slides


def test_sample_theme_on_dark_slides_and_dull_accents(tmp_path):
    frames = [_frame(tmp_path / "d.png", (0, 8, 16), (20, 50, 80))]  # dull navy
    theme = sample_theme(frames)
    assert theme is not None
    assert theme.fg == (240, 240, 240)
    assert theme.accent != (16, 48, 80)  # too close to the background: default used


def test_sample_theme_none_without_usable_frames(tmp_path):
    assert sample_theme([]) is None
    bad = tmp_path / "bad.png"
    bad.write_bytes(b"not an image")
    assert sample_theme([bad]) is None


def test_plain_theme_follows_settings():
    theme = plain_theme(DEFAULT_CONFIG)
    assert theme.bg == (0, 0, 0) and theme.fg == (240, 240, 240)


def test_neighbour_frames_skip_replaced_slides():
    frames = [Path(f"{i}.png") for i in range(5)]
    assert neighbour_frames(2, frames, {2}) == [Path("1.png"), Path("3.png")]
    assert neighbour_frames(2, frames, {1, 2, 3}) == [Path("0.png"), Path("4.png")]
    assert neighbour_frames(0, frames, {0}) == [Path("1.png")]
    assert neighbour_frames(1, frames[:2], {0, 1}) == []


def test_choose_code_slides_modes():
    assert choose_code_slides(DECK, [1], "auto") == [1]
    assert choose_code_slides(DECK, [0, 1], "auto") == [1]  # slide 0 has no code
    assert choose_code_slides(DECK, [1], "always") == [1, 2]
    assert choose_code_slides(DECK, [1], "never") == []


def test_render_code_slide_draws_theme_and_code(tmp_path):
    theme = theme_from((248, 248, 248), (150, 60, 230))
    out = render_code_slide(tmp_path / "c.png", CODE, theme, (960, 540))
    with Image.open(out) as im:
        assert im.size == (960, 540)
        assert im.getpixel((5, 5)) == (248, 248, 248)  # template background
        colours = set(im.getdata())
    assert theme.panel in colours  # the code box
    assert len(colours) > 20  # highlighted text, title, bullets


def test_render_code_slide_shrinks_long_code_to_fit(tmp_path):
    long = parse_markdown("# Long\n\n```python\n" + "\n".join(f"x{i} = {i}" for i in range(60)) + "\n```\n")[0]
    out = render_code_slide(tmp_path / "l.png", long, theme_from((0, 0, 0)), (960, 540))
    with Image.open(out) as im:
        # The code box ends above the bottom margin.
        assert im.getpixel((480, 530)) == (0, 0, 0)


def _texts(slide: Any) -> list[str]:
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]


def _shape_with(slide: Any, needle: str) -> Any:
    return next(sh for sh in slide.shapes if sh.has_text_frame and needle in sh.text_frame.text)


def _texts_of(shape: Any) -> str | None:
    return shape.text_frame.text if shape.has_text_frame else None


def _notes(slide: Any) -> str:
    return slide.notes_slide.notes_text_frame.text


def _designed_deck(path: Path) -> Path:
    prs = Presentation()
    for title in ("Intro", "Prose about code", "More"):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = title
        run.font.size = Pt(40)
        run.font.name = "Poppins"
        slide.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1)).text_frame.text = "body"
        notes: Any = slide.notes_slide.notes_text_frame
        notes.text = f"Narration for {title}"
    prs.save(str(path))
    return path


def test_rewrite_pptx_replaces_only_chosen_slides(tmp_path):
    deck = _designed_deck(tmp_path / "d.pptx")
    rewrite_pptx_code_slides(deck, DECK, [1], theme_from((248, 248, 248), (150, 60, 230)))
    prs = Presentation(str(deck))
    assert len(prs.slides) == 3
    assert [_notes(s) for s in prs.slides] == [
        "Narration for Intro", "Narration for Prose about code", "Narration for More"]
    assert _texts(prs.slides[0]) == ["Intro", "body"]

    texts = _texts(prs.slides[1])
    assert "Code" in texts
    assert "def f(x):\n    return x + 1" in texts
    assert "• adds one" in texts
    assert "Prose about code" not in texts and "body" not in texts


def test_rewrite_pptx_code_is_coloured_monospace_and_title_borrows_font(tmp_path):
    deck = _designed_deck(tmp_path / "d.pptx")
    rewrite_pptx_code_slides(deck, DECK, [1], theme_from((0, 8, 16)))
    slide = Presentation(str(deck)).slides[1]
    code_box = _shape_with(slide, "def f")
    runs = [r for p in code_box.text_frame.paragraphs for r in p.runs if r.text.strip()]
    assert {r.font.name for r in runs} == {"Courier New"}
    assert len({str(r.font.color.rgb) for r in runs}) >= 3
    title_run = _shape_with(slide, "Code").text_frame.paragraphs[0].runs[0]
    assert title_run.font.name == "Poppins" and title_run.font.size == Pt(40)


def test_rewrite_pptx_with_no_indices_leaves_file_untouched(tmp_path):
    deck = _designed_deck(tmp_path / "d.pptx")
    before = deck.read_bytes()
    rewrite_pptx_code_slides(deck, DECK, [], theme_from((0, 0, 0)))
    assert deck.read_bytes() == before


def test_sample_theme_background_is_exact_not_rounded(tmp_path):
    theme = sample_theme([_frame(tmp_path / "w.png", (250, 250, 251))])
    assert theme is not None and theme.bg == (250, 250, 251)
