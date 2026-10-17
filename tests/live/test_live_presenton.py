"""Live Presenton contract (ADR-0009, ADR-0013).

Needs a running Presenton (tested: v0.9.11-beta) at providers.presenton.base_url
or PRESENTON_URL, and PRESENTON_API_KEY when it requires login. Skips when
no server answers. Uses the server's own LLM, so it costs a few cents with a
cloud model. Run with:
    uv run pytest tests/live/test_live_presenton.py --run-live -q
"""

import pytest
import requests

from slide_stream.fidelity import pdf_slide_texts
from slide_stream.parser import parse_markdown
from slide_stream.presenton import (
    PresentonClient,
    inject_speaker_notes,
    resolve_presenton_settings,
    review_designed_deck,
)
from slide_stream.slide_render import pdf_to_slide_images

DECK = """# Hash Tables

- Find a value without searching
- Week 6, Data Structures

# A Hash Function

```python
def bucket_for(key: str, n: int) -> int:
    return hash(key) % n
```

- hash() turns a key into an integer
"""


@pytest.fixture
def settings(live_config):
    s = resolve_presenton_settings(live_config)
    try:
        r = requests.get(f"{s['base_url']}/api/v1/auth/status", timeout=5)
    except requests.RequestException:
        pytest.skip(f"no Presenton at {s['base_url']}")
    if r.status_code != 200:
        pytest.skip(f"Presenton health check answered {r.status_code}")
    return s


def test_live_presenton_designs_aligned_deck_and_rebuilds_code(settings, live_config, tmp_path):
    slides = parse_markdown(DECK)
    client = PresentonClient(
        settings["base_url"], api_key=settings["api_key"],
        instructions=settings["instructions"],
    )
    pptx, pdf = tmp_path / "deck.pptx", tmp_path / "deck.pdf"
    result = client.generate_deck_pair(slides, str(settings["template"]), pptx, pdf)

    # Contract: one designed slide per source slide, both formats downloaded.
    assert result.presentation_id
    assert pptx.stat().st_size > 10_000 and pdf.stat().st_size > 1_000
    frames = pdf_to_slide_images(pdf, tmp_path / "frames", width=640)
    assert len(frames) == len(slides) == len(pdf_slide_texts(pdf))

    # ADR-0013: the code always ends up on the slide, in both outputs.
    review = review_designed_deck(
        slides, {**settings, "code_slides": "auto"}, live_config,
        pptx_path=pptx, pdf_path=pdf, frames=frames, frame_size=(640, 360),
        work_dir=tmp_path / "code",
    )
    from slide_stream.fidelity import pptx_slide_texts

    assert "return hash(key) % n" in pptx_slide_texts(pptx)[1].replace("\u00a0", " ")
    if review.rebuilt:
        assert review.rebuilt == [1]
        assert frames[1].name == "code_slide_2.png"

    inject_speaker_notes(pptx, ["Intro narration.", "Code narration."])
    from pptx import Presentation

    notes = [
        str(getattr(s.notes_slide.notes_text_frame, "text", ""))
        for s in Presentation(str(pptx)).slides
    ]
    assert notes == ["Intro narration.", "Code narration."]
    print(f"\nPresenton rebuilt slides: {[i + 1 for i in review.rebuilt]}")
