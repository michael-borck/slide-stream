"""Tests for the designed-deck fidelity check (ADR-0013)."""

import json

from pptx import Presentation
from pptx.util import Inches

from slide_stream.fidelity import (
    check_fidelity,
    check_slide,
    designed_texts,
    pptx_slide_texts,
)
from slide_stream.parser import parse_markdown

CODE_SLIDE = parse_markdown(
    "# A Hash Function\n\n"
    "```python\ndef bucket_for(key, n):\n    return hash(key) % n\n```\n\n"
    "- `hash()` turns a value into an integer\n"
)[0]


def test_code_kept_only_when_every_line_is_shown():
    shown = "A Hash Function\ndef bucket_for(key, n):\n  return hash(key) % n\nhash() turns a value into an integer"
    assert check_slide(1, CODE_SLIDE, shown).code_kept is True
    # Quoting one line in prose is not showing the code.
    prose = "A Hash Function\nThe line `return hash(key) % n` folds the hash into range."
    result = check_slide(1, CODE_SLIDE, prose)
    assert result.code_kept is False
    assert result.code_lines_kept == 0.5
    assert "code missing" in result.problems()[0]


def test_slides_without_code_or_tables_report_none():
    slide = parse_markdown("# Plain\n\n- just text\n")[0]
    result = check_slide(1, slide, "Plain\njust text")
    assert result.code_kept is None and result.table_kept is None
    assert result.problems() == []


def test_title_and_content_kept():
    slide = parse_markdown("# Searching a List\n\n- one million student records\n- check each record\n")[0]
    same = check_slide(1, slide, "Searching a List\nOne million student records. Check each record.")
    assert same.title_kept and same.content_kept == 1.0
    changed = check_slide(1, slide, "Linear Search\nStudent records")
    assert not changed.title_kept
    assert changed.content_kept < 0.5
    assert "title changed" in changed.problems()


def test_table_cells_and_invented_row():
    slide = parse_markdown("# Speed\n\n| Op | List |\n|---|---|\n| Insert | O(1) |\n")[0]
    designed = "Speed\nOp List\nInsert O(1)\nSummary Slower as data grows"  # PDF text: no pipes
    result = check_slide(1, slide, designed)
    assert result.table_kept == 1.0
    assert "Summary Slower" in result.invented


def test_invented_people_and_numbers_but_not_step_numbers():
    deck = parse_markdown("# Every Cup Starts on a Mountain\n\n- 1,800 metres up\n")
    designed = "Every Cup Starts on a Mountain\n01 1,800 metres up\nMaria Lopez\nFarm Owner\nSince 1987"
    report = check_fidelity(deck, [designed])
    invented = report.slides[0].invented
    assert "Maria Lopez" in invented and "1987" in invented
    assert "1,800" not in invented and "01" not in invented


def test_names_found_elsewhere_in_the_deck_are_not_invented():
    deck = parse_markdown("# One\n\n- hash tables\n\n# Two\n\n- open addressing\n")
    report = check_fidelity(deck, ["One\nHash Tables", "Two\nOpen Addressing"])
    assert all(not s.invented for s in report.slides)


def test_report_lists_slides_needing_code_and_writes_json(tmp_path):
    deck = [CODE_SLIDE, parse_markdown("# Plain\n\n- text\n")[0]]
    report = check_fidelity(deck, ["A Hash Function\nprose only", "Plain\ntext"], min_kept=0.5)
    assert report.code_missing() == [0]
    out = report.write_json(tmp_path / "deck.fidelity.json")
    data = json.loads(out.read_text())
    assert data["min_kept"] == 0.5
    assert data["slides"][0]["code_kept"] is False
    assert data["slides"][0]["problems"]
    assert data["slides"][1]["problems"] == []


def test_pptx_slide_texts_reads_groups_and_tables(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "Title"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1)).text_frame.text = "In a group"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(1, 1).text = "Op", "O(1)"
    path = tmp_path / "d.pptx"
    prs.save(str(path))
    (text,) = pptx_slide_texts(path)
    assert "Title" in text and "In a group" in text and "Op" in text and "O(1)" in text


def test_designed_texts_prefers_pdf_and_falls_back_to_pptx(tmp_path):
    import pytest

    pdfium = pytest.importorskip("pypdfium2")
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "From pptx"
    pptx = tmp_path / "d.pptx"
    prs.save(str(pptx))
    assert designed_texts(pptx, tmp_path / "missing.pdf") == ["From pptx"]
    doc = pdfium.PdfDocument.new()
    doc.new_page(100, 100)
    pdf = tmp_path / "d.pdf"
    doc.save(str(pdf))
    doc.close()
    assert designed_texts(pptx, pdf) == [""]  # the (blank) PDF wins
