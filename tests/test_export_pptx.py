"""Tests for the .pptx exporter (slides -> editable PowerPoint deck)."""

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from slide_stream.export_pptx import export_pptx


def _shape_texts(slide):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            texts.append(shape.text_frame.text)
    return "\n".join(texts)


def _picture_count(slide):
    return sum(
        1 for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE
    )


def test_exports_titles_bullets_and_notes(tmp_path):
    out = tmp_path / "deck.pptx"
    result = export_pptx(
        [
            {
                "title": "Coffee",
                "content": ["It fuels developers", "Espresso or filter"],
                "notes": "Coffee keeps the lights on.",
            },
            {
                "title": "Tea",
                "content": ["The quieter alternative"],
                "notes": "Tea is for contemplation.",
            },
        ],
        out,
    )
    assert result == out and out.exists()

    prs = Presentation(str(out))
    assert len(prs.slides) == 2
    first_texts = _shape_texts(prs.slides[0])
    assert "Coffee" in first_texts
    assert "It fuels developers" in first_texts
    assert "Espresso or filter" in first_texts
    notes = prs.slides[0].notes_slide.notes_text_frame
    assert notes is not None
    assert notes.text == "Coffee keeps the lights on."
    assert "Tea" in _shape_texts(prs.slides[1])


def test_embeds_artwork_and_fits_the_slide(tmp_path):
    img = tmp_path / "art.png"
    Image.new("RGB", (1920, 1080), "steelblue").save(img)
    out = tmp_path / "deck.pptx"
    export_pptx(
        [{"title": "T", "content": ["b"], "image_path": str(img), "notes": "n"}],
        out,
    )

    prs = Presentation(str(out))
    slide = prs.slides[0]
    assert _picture_count(slide) == 1
    pic = next(
        s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    # Picture stays inside the slide bounds.
    width, height = prs.slide_width, prs.slide_height
    assert width is not None and height is not None
    assert pic.left >= 0 and pic.top >= 0
    assert pic.left + pic.width <= width
    assert pic.top + pic.height <= height


def test_tall_image_is_scaled_to_fit(tmp_path):
    img = tmp_path / "tall.png"
    Image.new("RGB", (400, 2000), "salmon").save(img)
    out = tmp_path / "deck.pptx"
    export_pptx([{"title": "T", "content": [], "image_path": str(img)}], out)

    prs = Presentation(str(out))
    pic = next(
        s for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    width, height = prs.slide_width, prs.slide_height
    assert width is not None and height is not None
    assert pic.height <= height
    assert pic.width <= width


def test_missing_image_file_renders_text_only(tmp_path):
    out = tmp_path / "deck.pptx"
    export_pptx(
        [{"title": "T", "content": ["b"], "image_path": str(tmp_path / "nope.png")}],
        out,
    )
    prs = Presentation(str(out))
    assert _picture_count(prs.slides[0]) == 0
    assert "b" in _shape_texts(prs.slides[0])


def test_slide_without_title_gets_placeholder(tmp_path):
    out = tmp_path / "deck.pptx"
    export_pptx([{"title": "", "content": [], "notes": ""}], out)
    prs = Presentation(str(out))
    assert "Slide" in _shape_texts(prs.slides[0])


def test_creates_missing_parent_directories(tmp_path):
    out = tmp_path / "a" / "b" / "deck.pptx"
    result = export_pptx([{"title": "T", "content": []}], out)
    assert result.exists()


@pytest.mark.parametrize("empty_notes", [None, "", "   "])
def test_blank_notes_are_skipped(tmp_path, empty_notes):
    out = tmp_path / "deck.pptx"
    export_pptx([{"title": "T", "content": [], "notes": empty_notes}], out)
    prs = Presentation(str(out))
    slide = prs.slides[0]
    # Accessing notes_slide creates it lazily; the text must stay blank.
    notes = slide.notes_slide.notes_text_frame
    assert notes is not None
    assert notes.text.strip() == ""
