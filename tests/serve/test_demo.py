"""Web UI hosted demo: teaser, topic drafting, caps, rate limits, hardening."""

import copy
from io import BytesIO

import pytest
import yaml

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient  # noqa: E402

from slide_stream import serve  # noqa: E402
from slide_stream.config_loader import DEFAULT_CONFIG  # noqa: E402


def _tiny_png() -> bytes:
    """A real (decodable) 4x4 PNG for photo uploads."""
    from PIL import Image

    buf = BytesIO()
    Image.new("RGB", (4, 4), "red").save(buf, format="PNG")
    return buf.getvalue()


def _llm_config():
    cfg = copy.deepcopy(DEFAULT_CONFIG)
    cfg["providers"]["llm"]["provider"] = "openai"
    return cfg


def test_demo_mode_needs_no_token(base_config):
    """Demo mode is open even when a token is configured."""
    client = TestClient(serve.create_app(config=base_config, token="secret", demo=True))
    cfg = client.get("/api/config").json()
    assert cfg["auth_required"] is False
    assert cfg["demo"] is True
    assert cfg["limits"] == {
        "max_slides": serve.DEMO_MAX_SLIDES,
        "jobs_per_hour": serve.DEMO_JOBS_PER_HOUR,
        "job_ttl_minutes": serve.JOB_TTL_SECONDS // 60,
    }
    # protected endpoint reachable without a token
    assert client.get("/api/jobs/whatever").status_code == 404


# --- teaser: text formats, voice picker, avatar art ----------------------------


def test_build_topic_prompt_contains_topic_and_count():
    from slide_stream.draft import build_topic_prompt

    prompt = build_topic_prompt("Why sleep matters", 5)
    assert "Why sleep matters" in prompt
    assert "exactly 5 slides" in prompt
    assert "'# ' heading" in prompt  # output contract for the parser


def test_topic_draft_endpoint(base_config, monkeypatch):
    captured: dict[str, object] = {}

    def fake_draft(topic, slides, provider, model, base_url,
                   api_key=None, think=None):
        captured.update(topic=topic, slides=slides)
        return "# A\n\n- x\n\n# B\n\n- y\n"

    monkeypatch.setattr(serve, "_do_topic_draft", fake_draft)
    client = TestClient(serve.create_app(config=_llm_config(), token=None))
    r = client.post(
        "/api/draft-deck",
        json={"topic": "Why sleep matters for learning", "slides": 2},
    )
    assert r.status_code == 200, r.text
    assert "# A" in r.json()["markdown"]
    assert captured["topic"] == "Why sleep matters for learning"
    assert captured["slides"] == 2


def test_topic_draft_demo_caps_slides_and_defaults_to_five(monkeypatch):
    captured: dict[str, object] = {}

    def fake_draft(topic, slides, provider, model, base_url,
                   api_key=None, think=None):
        captured["slides"] = slides
        return "# A\n\n- x\n"

    monkeypatch.setattr(serve, "_do_topic_draft", fake_draft)
    client = TestClient(serve.create_app(config=_llm_config(), demo=True))
    r = client.post("/api/draft-deck", json={"topic": "dogs", "slides": 9})
    assert r.status_code == 200
    assert r.json()["slides"] == serve.DEMO_MAX_SLIDES
    assert captured["slides"] == serve.DEMO_MAX_SLIDES


def test_topic_draft_requires_topic(base_config):
    client = TestClient(serve.create_app(config=_llm_config(), token=None))
    assert client.post("/api/draft-deck", json={"topic": "  "}).status_code == 400


def test_topic_draft_requires_llm():
    client = TestClient(serve.create_app(config=copy.deepcopy(DEFAULT_CONFIG), token=None))
    r = client.post("/api/draft-deck", json={"topic": "dogs at the park"})
    assert r.status_code == 400
    assert "LLM" in r.json()["detail"]


def test_topic_draft_demo_rate_limited(monkeypatch):
    monkeypatch.setattr(
        serve, "_do_topic_draft", lambda *a, **k: "# A\n\n- x\n"
    )
    client = TestClient(serve.create_app(config=_llm_config(), demo=True))
    for _ in range(serve.DEMO_DRAFTS_PER_HOUR):
        assert client.post("/api/draft-deck", json={"topic": "cats"}).status_code == 200
    over = client.post("/api/draft-deck", json={"topic": "cats"})
    assert over.status_code == 429
    assert "drafts per hour" in over.json()["detail"]


def test_demo_accepts_txt_with_slides(base_config, mocker, monkeypatch):
    """A plain .txt structured with '# ' headings is parsed as a deck."""
    seen: dict[str, str] = {}

    def fake_pipeline(job, deck_path, job_yaml, voice_path, photo_path):
        assert job.workdir is not None
        out = job.workdir / "output.mp4"
        out.write_bytes(b"V")
        job.status = "done"
        job.output_path = out
        seen["deck_name"] = deck_path.name

    mocker.patch.object(serve, "_run_video_pipeline", side_effect=fake_pipeline)
    monkeypatch.setattr(
        serve.ThreadPoolExecutor, "submit", lambda self, fn, *a, **k: fn(*a, **k)
    )
    client = TestClient(serve.create_app(config=base_config, demo=True))
    r = client.post("/api/jobs", files={
        "deck": ("notes.txt", b"# First slide\n- a\n\n# Second\n- b\n", "text/plain"),
    })
    assert r.status_code == 200, r.text
    assert seen["deck_name"] == "deck.md"  # stored under the canonical name


def test_demo_rejects_text_without_slide_structure(base_config, mocker):
    client = TestClient(serve.create_app(config=base_config, demo=True))
    r = client.post("/api/jobs", files={
        "deck": ("notes.txt", b"just some prose,\nno headings at all\n", "text/plain"),
    })
    assert r.status_code == 400
    assert "No slides found" in r.json()["detail"]
    assert "'# '" in r.json()["detail"]


def test_demo_accepts_qmd(base_config, mocker, monkeypatch):
    """A Quarto (.qmd) file with front matter parses; front matter is dropped."""
    seen: dict[str, int] = {}

    def fake_pipeline(job, deck_path, job_yaml, voice_path, photo_path):
        assert job.workdir is not None
        out = job.workdir / "output.mp4"
        out.write_bytes(b"V")
        job.status = "done"
        job.output_path = out
        seen["slides"] = len(serve._parse_deck_slides(deck_path))

    mocker.patch.object(serve, "_run_video_pipeline", side_effect=fake_pipeline)
    monkeypatch.setattr(
        serve.ThreadPoolExecutor, "submit", lambda self, fn, *a, **k: fn(*a, **k)
    )
    client = TestClient(serve.create_app(config=base_config, demo=True))
    qmd = b"---\ntitle: Deck\n---\n\n# One\n\n- a\n\n# Two\n\n- b\n"
    r = client.post("/api/jobs", files={"deck": ("deck.qmd", qmd, "text/markdown")})
    assert r.status_code == 200, r.text
    assert seen["slides"] == 2


def test_voice_name_maps_to_configured_stock_voice(base_config, tmp_path):
    """A teaser voice choice overrides stored voice keys and resolves the
    display name back to the configured file ('Emily' -> 'Emily.wav')."""
    base = copy.deepcopy(base_config)
    base["providers"]["tts"].update({
        "provider": "voicebox",
        "profile_id": "stored-id",
        "voice": "old-stock.wav",
        "reference_text": "stale",
        "voice_choices": ["Emily.wav", "James.wav"],
    })
    workdir = tmp_path / "job"
    workdir.mkdir()
    cfg = yaml.safe_load(serve._build_job_config(
        base, workdir, {"avatar": True, "voice_name": "Emily"}, None, None
    ).read_text())
    tts = cfg["providers"]["tts"]
    assert tts["voice"] == "Emily.wav"
    assert "profile_id" not in tts and "reference_text" not in tts


def test_api_config_exposes_voices(base_config):
    base_config["providers"]["tts"]["voice_choices"] = ["Emily.wav", "James.wav"]
    client = TestClient(serve.create_app(config=base_config, demo=True))
    assert client.get("/api/config").json()["voices"] == ["Emily", "James"]


def test_avatar_image_endpoint(base_config):
    from slide_stream.avatars import avatar_names

    client = TestClient(serve.create_app(config=base_config, demo=True))
    name = avatar_names()[0]
    ok = client.get(f"/api/avatars/{name}")
    assert ok.status_code == 200
    assert ok.headers["content-type"].startswith("image/")
    assert client.get("/api/avatars/not-a-mascot").status_code == 404


def test_demo_defaults_presenter_to_first_and_last(base_config, mocker):
    """Teaser economics: without explicit placement, demo presenters appear
    on the opening and closing slides only (wan-s2v is minutes per clip)."""
    real = serve._build_job_config
    captured: dict[str, object] = {}

    def spy(base, workdir, options, *args, **kwargs):
        captured["slides"] = options.get("avatar_slides")
        return real(base, workdir, options, *args, **kwargs)

    mocker.patch.object(serve, "_build_job_config", side_effect=spy)
    mocker.patch.object(serve, "_run_job")
    client = TestClient(serve.create_app(config=base_config, demo=True))
    r = client.post("/api/jobs", files={"deck": ("deck.md", b"# One\n- a\n")},
                    data={"avatar_name": "teddy"})
    assert r.status_code == 200, r.text
    assert captured["slides"] == "first,last"

    # Explicit placement always wins.
    captured.clear()
    r = client.post("/api/jobs", files={"deck": ("deck.md", b"# One\n- a\n")},
                    data={"avatar_name": "teddy", "avatar_slides": "first"})
    assert r.status_code == 200
    assert captured["slides"] == "first"


def test_demo_drops_voice_and_photo_uploads(base_config, mocker, monkeypatch):
    """The teaser accepts no biometrics: voice/photo uploads are never
    written to disk or passed to the render, even from API clients."""
    seen: dict[str, object] = {}

    def fake_pipeline(job, deck_path, job_yaml, voice_path, photo_path):
        assert job.workdir is not None
        out = job.workdir / "output.mp4"
        out.write_bytes(b"V")
        job.status = "done"
        job.output_path = out
        seen["voice"] = voice_path
        seen["photo"] = photo_path

    mocker.patch.object(serve, "_run_video_pipeline", side_effect=fake_pipeline)
    monkeypatch.setattr(
        serve.ThreadPoolExecutor, "submit", lambda self, fn, *a, **k: fn(*a, **k)
    )
    client = TestClient(serve.create_app(config=base_config, demo=True))
    r = client.post("/api/jobs", files={
        "deck": ("deck.md", b"# One\n- a\n"),
        "voice": ("me.wav", b"AUDIO", "audio/wav"),
        "photo": ("me.png", _tiny_png(), "image/png"),
    })
    assert r.status_code == 200
    assert seen["voice"] is None
    assert seen["photo"] is None
    # Nothing biometric ever landed in the job workdir.
    workdir = serve._JOBS[r.json()["job_id"]].workdir
    assert workdir is not None
    assert not list(workdir.glob("voice*"))
    assert not list(workdir.glob("photo*"))


def test_demo_slide_cap_trims_to_five(base_config, mocker, monkeypatch):
    """A big deck is trimmed to the first 5 slides before rendering."""
    captured: dict[str, object] = {}

    def fake_pipeline(job, deck_path, job_yaml, voice_path, photo_path):
        assert job.workdir is not None
        out = job.workdir / "output.mp4"
        out.write_bytes(b"V")
        job.status = "done"
        job.output_path = out
        captured["slides"] = len(serve._parse_deck_slides(deck_path))

    mocker.patch.object(serve, "_run_video_pipeline", side_effect=fake_pipeline)
    monkeypatch.setattr(
        serve.ThreadPoolExecutor, "submit", lambda self, fn, *a, **k: fn(*a, **k)
    )
    client = TestClient(serve.create_app(config=base_config, demo=True))
    big_deck = "".join(f"# Slide {i}\n- point\n" for i in range(10))
    r = client.post("/api/jobs", files={"deck": ("deck.md", big_deck.encode())})
    assert r.status_code == 200
    body = r.json()
    assert "first 5" in body["notice"]
    assert "desktop app" in body["notice"]
    assert captured["slides"] == serve.DEMO_MAX_SLIDES

    # A small deck is accepted untouched, with no notice.
    ok = client.post("/api/jobs", files={"deck": ("deck.md", b"# One\n- a\n")})
    assert ok.status_code == 200
    assert ok.json()["notice"] is None


def test_demo_pptx_cap_trims_to_five(base_config, mocker, monkeypatch):
    """An oversized .pptx is trimmed to DEMO_MAX_SLIDES slides in place."""
    from pptx import Presentation

    mocker.patch.object(serve, "_run_job")
    monkeypatch.setattr(
        serve.ThreadPoolExecutor, "submit", lambda self, fn, *a, **k: fn(*a, **k)
    )
    prs = Presentation()
    for i in range(8):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Slide {i}"  # type: ignore[union-attr]
    buf = BytesIO()
    prs.save(buf)

    client = TestClient(serve.create_app(config=base_config, demo=True))
    r = client.post(
        "/api/jobs",
        files={"deck": ("deck.pptx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument."
                        "presentationml.presentation")},
        data={"output": "pptx"},
    )
    assert r.status_code == 200, r.text
    assert "first 5" in r.json()["notice"]
    # The deck handed to the render has exactly DEMO_MAX_SLIDES slides.
    jobs = [j for j in serve._JOBS.values() if j.workdir]
    assert jobs
    workdir = jobs[-1].workdir
    assert workdir is not None
    pptx_files = list(workdir.glob("*.pptx"))
    assert pptx_files
    assert serve._count_slides(pptx_files[0]) == serve.DEMO_MAX_SLIDES


def test_truncate_markdown_deck_heading_style():
    from slide_stream.parser import parse_markdown

    md = "".join(f"# S{i}\n\n- point {i}\n\n" for i in range(8))
    out = serve._truncate_markdown_deck(md, 3)
    assert len(parse_markdown(out)) == 3
    assert "# S0" in out and "# S2" in out and "# S3" not in out


def test_truncate_markdown_deck_separator_style():
    from slide_stream.parser import parse_markdown

    blocks = [f"# S{i}\npoint {i}" for i in range(6)]
    md = "\n---\n".join(blocks) + "\n"
    out = serve._truncate_markdown_deck(md, 2)
    slides = parse_markdown(out)
    assert len(slides) == 2
    assert slides[0]["title"] == "S0" and slides[1]["title"] == "S1"


def test_truncate_markdown_deck_keeps_front_matter():
    from slide_stream.parser import parse_markdown

    md = "---\ntitle: Deck\ntheme: dark\n---\n# A\n\n# B\n\n# C\n"
    out = serve._truncate_markdown_deck(md, 2)
    assert out.startswith("---\ntitle: Deck")
    assert len(parse_markdown(out)) == 2


def test_truncate_markdown_deck_matches_parser_on_fences():
    """Boundary rules intentionally mirror parse_markdown — a '#' line inside
    an unclosed-looking fence counts (and cuts) the same way in both, so a
    trimmed deck never parses to more slides than were kept."""
    from slide_stream.parser import parse_markdown

    md = "# A\n```text\n# fake\n```\n\n# B\n\n# C\n"
    out = serve._truncate_markdown_deck(md, 3)
    assert "# B" in out and "# C" not in out
    assert len(parse_markdown(out)) <= 3


def test_truncate_markdown_deck_short_deck_unchanged():
    from slide_stream.parser import parse_markdown

    md = "# A\n\n# B\n"
    out = serve._truncate_markdown_deck(md, 5)
    assert len(parse_markdown(out)) == 2


def test_demo_rate_limit(base_config, mocker):
    client = TestClient(serve.create_app(config=base_config, demo=True))
    mocker.patch.object(serve, "_run_job")
    deck = {"deck": ("deck.md", b"# One\n- a\n")}
    for _ in range(serve.DEMO_JOBS_PER_HOUR):
        assert client.post("/api/jobs", files=deck).status_code == 200
    over = client.post("/api/jobs", files=deck)
    assert over.status_code == 429
    assert "per hour" in over.json()["detail"]


def test_demo_rate_window_expires():
    now = 1_000_000.0
    ip = "1.2.3.4"
    for _ in range(serve.DEMO_JOBS_PER_HOUR):
        assert serve._demo_rate_ok(ip, now=now)
    assert not serve._demo_rate_ok(ip, now=now + 10)
    # an hour later the window has rolled over
    assert serve._demo_rate_ok(ip, now=now + 3601)


def test_non_demo_still_requires_token(base_config):
    client = TestClient(serve.create_app(config=base_config, token="secret", demo=False))
    assert client.get("/api/jobs/whatever").status_code == 401


# --- local (desktop) mode ------------------------------------------------------


def test_demo_unparseable_deck_fails_closed(base_config, mocker):
    mocker.patch.object(serve, "_run_job")
    mocker.patch.object(serve, "_count_slides", return_value=None)
    client = TestClient(serve.create_app(config=base_config, demo=True))
    r = client.post("/api/jobs", files={"deck": ("deck.md", b"# One\n- a\n")})
    assert r.status_code == 400
    assert "parse" in r.json()["detail"].lower()


def test_demo_status_hides_log_and_token(base_config):
    client = TestClient(serve.create_app(config=base_config, demo=True))
    serve._JOBS["j1"] = serve.Job(id="j1", status="running",
                                  log="Traceback: /secret/server/path",
                                  download_token="dltok")
    s = client.get("/api/jobs/j1").json()
    assert s["status"] == "running"  # coarse progress the UI shows
    assert s["log"] == ""
    assert "token" not in s  # else anyone with the UUID could download


def test_status_returns_log_and_token_when_authed(base_config):
    client = TestClient(serve.create_app(config=base_config, token="secret"))
    serve._JOBS["j1"] = serve.Job(id="j1", status="running",
                                  log="rendering slide 2",
                                  download_token="dltok")
    s = client.get("/api/jobs/j1",
                   headers={"Authorization": "Bearer secret"}).json()
    assert s["log"] == "rendering slide 2"
    assert s["token"] == "dltok"


def test_demo_download_needs_job_token_and_wipes_workdir(base_config, tmp_path):
    client = TestClient(serve.create_app(config=base_config, demo=True))
    workdir = tmp_path / "job"
    workdir.mkdir()
    out = workdir / "output.mp4"
    out.write_bytes(b"VID")
    serve._JOBS["j1"] = serve.Job(id="j1", status="done", workdir=workdir,
                                  output_path=out, download_token="dltok")
    # The UUID alone is no longer enough in demo mode.
    assert client.get("/api/jobs/j1/result").status_code == 401
    r = client.get("/api/jobs/j1/result", params={"t": "dltok"})
    assert r.status_code == 200
    assert r.content == b"VID"
    # "Nothing stored": one download, then the whole job is gone.
    assert not workdir.exists()
    assert "j1" not in serve._JOBS


# --- local (desktop) mode CSRF/DNS-rebinding guards ---------------------------
