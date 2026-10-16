"""Fidelity check for designed decks (ADR-0013).

A design backend such as Presenton rewrites slide text to fit its layouts.
SlideStream narrates from the source, so anything the designer drops or
invents makes voice and picture disagree. This module compares each
designed slide's text with its source slide and reports what changed.

The checks are heuristics on text. They flag slides for a person to look
at; they do not judge design.
"""

import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

DEFAULT_MIN_KEPT = 0.5
FIDELITY_MODES = ("warn", "strict", "off")

_WORD_RE = re.compile(r"[a-z0-9]+(?:'[a-z]+)?")
_STOPWORDS = {
    "a", "an", "and", "are", "as", "at", "be", "but", "by", "can", "do",
    "does", "for", "from", "has", "have", "how", "if", "in", "into", "is",
    "it", "its", "of", "on", "or", "so", "than", "that", "the", "their",
    "then", "there", "this", "to", "was", "we", "what", "when", "which",
    "with", "you", "your",
}
# Two or more capitalised words in a row: a likely name or role.
_NAME_RE = re.compile(r"\b[A-Z][a-z]+(?:[ \t]+[A-Z][a-z]*\.?)+")  # one line only
# Numbers worth checking: 3+ digits (1,800 / 2026), decimals, percentages.
# Small integers are skipped because layouts number their steps 1, 2, 3.
_NUMBER_RE = re.compile(r"\d[\d,]*\.\d+|\d{1,3}(?:,\d{3})+|\d{3,}|\d+(?:\.\d+)?%")
_MAX_INVENTED = 5


def _words(text: str) -> list[str]:
    return _WORD_RE.findall(text.lower())


def _content_words(text: str) -> set[str]:
    return {w for w in _words(text) if len(w) > 2 and w not in _STOPWORDS}


def _squash(text: str) -> str:
    """Lowercase with all whitespace removed, for verbatim code matching."""
    return re.sub(r"\s+", "", text.lower())


def _norm_number(n: str) -> str:
    return n.replace(",", "").rstrip("%")


@dataclass
class SlideFidelity:
    """How faithfully one designed slide reproduces its source."""

    index: int  # 1-based slide number
    title: str
    title_kept: bool
    content_kept: float  # share of the source's content words present, 0..1
    code_kept: bool | None = None  # None when the source slide has no code
    code_lines_kept: float | None = None
    table_kept: float | None = None  # share of source table cells present
    invented: list[str] = field(default_factory=list)

    def problems(self, min_kept: float = DEFAULT_MIN_KEPT) -> list[str]:
        """Short descriptions of what looks wrong, for the report."""
        out: list[str] = []
        if self.code_kept is False:
            out.append(f"code missing ({self.code_lines_kept or 0:.0%} of lines shown)")
        if not self.title_kept:
            out.append("title changed")
        if self.content_kept < min_kept:
            out.append(f"only {self.content_kept:.0%} of the text kept")
        if self.table_kept is not None and self.table_kept < 1.0:
            out.append(f"{self.table_kept:.0%} of table cells kept")
        if self.invented:
            out.append("possibly invented: " + ", ".join(self.invented))
        return out


@dataclass
class FidelityReport:
    """Per-slide fidelity for a whole designed deck."""

    slides: list[SlideFidelity]
    min_kept: float = DEFAULT_MIN_KEPT

    def code_missing(self) -> list[int]:
        """0-based indices of slides whose source code is not on the slide."""
        return [s.index - 1 for s in self.slides if s.code_kept is False]

    def below_min_kept(self) -> list[int]:
        return [s.index - 1 for s in self.slides if s.content_kept < self.min_kept]

    def flagged(self) -> list[SlideFidelity]:
        return [s for s in self.slides if s.problems(self.min_kept)]

    def to_dict(self) -> dict[str, Any]:
        return {
            "min_kept": self.min_kept,
            "slides": [
                {**asdict(s), "problems": s.problems(self.min_kept)}
                for s in self.slides
            ],
        }

    def write_json(self, path: str | Path) -> Path:
        dest = Path(path)
        dest.write_text(json.dumps(self.to_dict(), indent=2), encoding="utf-8")
        return dest


def _source_text(slide: dict[str, Any]) -> str:
    """All of a source slide's visible text (title, content, notes excluded)."""
    parts = [str(slide.get("title", ""))]
    parts += [str(c) for c in slide.get("content", [])]
    return "\n".join(parts)


def _code_lines_kept(code_blocks: list[dict[str, Any]], designed: str) -> float:
    """Share of non-blank code lines that appear verbatim on the slide
    (whitespace ignored). Quoting identifiers in prose does not count."""
    squashed = _squash(designed)
    lines = [
        _squash(line)
        for block in code_blocks
        for line in str(block.get("code", "")).splitlines()
        if line.strip()
    ]
    if not lines:
        return 1.0
    return sum(1 for line in lines if line in squashed) / len(lines)


def _table_cells_kept(tables: list[list[list[str]]], designed: str) -> float:
    lowered = designed.lower()
    cells = [
        cell.strip().lower()
        for table in tables for row in table for cell in row if cell.strip()
    ]
    if not cells:
        return 1.0
    return sum(1 for c in cells if c in lowered) / len(cells)


def _invented(source: str, designed: str) -> list[str]:
    """Names and numbers on the designed slide that the source never had.

    ``source`` should be the whole deck's text: a designer that moves a
    phrase from one slide's wording into another's title has not invented
    it, but a person or figure found nowhere in the deck is suspicious.
    """
    source_vocab = set(_words(source))
    source_numbers = {_norm_number(n) for n in _NUMBER_RE.findall(source)}
    found: list[str] = []
    for m in _NAME_RE.finditer(designed):
        phrase = m.group(0).strip()
        words = _words(phrase)
        if words and not any(w in source_vocab for w in words):
            found.append(phrase)
    for n in _NUMBER_RE.findall(designed):
        if _norm_number(n) not in source_numbers:
            found.append(n)
    return list(dict.fromkeys(found))[:_MAX_INVENTED]


def check_slide(
    index: int, source: dict[str, Any], designed: str, deck_text: str | None = None
) -> SlideFidelity:
    """Compare one designed slide's text with its source slide.

    ``deck_text`` is the whole source deck's text, used for the "possibly
    invented" check; it defaults to this slide's own text.
    """
    title = str(source.get("title", "")).strip()
    has_title = bool(source.get("has_real_title", True)) and bool(title)
    title_kept = (
        not has_title
        or " ".join(_words(title)) in " ".join(_words(designed))
    )
    code_blocks = source.get("code_blocks") or []
    tables = source.get("tables") or []
    code_texts = {str(b.get("code", "")) for b in code_blocks}
    # Content words exclude code, which has its own verbatim check.
    prose = "\n".join(
        str(c) for c in source.get("content", []) if str(c) not in code_texts
    )
    source_words = _content_words(prose) or _content_words(title)
    designed_words = _content_words(designed)
    content_kept = (
        len(source_words & designed_words) / len(source_words) if source_words else 1.0
    )
    result = SlideFidelity(
        index=index, title=title, title_kept=title_kept,
        content_kept=round(content_kept, 3),
        invented=_invented(deck_text or _source_text(source), designed),
    )
    if code_blocks:
        kept = _code_lines_kept(code_blocks, designed)
        result.code_lines_kept = round(kept, 3)
        result.code_kept = kept >= 1.0
    if tables:
        result.table_kept = round(_table_cells_kept(tables, designed), 3)
    return result


def check_fidelity(
    source_slides: list[dict[str, Any]],
    designed_texts: list[str],
    min_kept: float = DEFAULT_MIN_KEPT,
) -> FidelityReport:
    """Compare every designed slide with its source, by position."""
    deck_text = "\n".join(_source_text(s) for s in source_slides)
    return FidelityReport(
        slides=[
            check_slide(i + 1, src, text, deck_text)
            for i, (src, text) in enumerate(zip(source_slides, designed_texts, strict=False))
        ],
        min_kept=min_kept,
    )


def designed_texts(pptx_path: str | Path | None, pdf_path: str | Path | None) -> list[str]:
    """What each designed slide shows, preferring the PDF.

    Presenton's .pptx export stores some elements, such as tables, as
    pictures, so the PDF (what the video shows) is the better source of
    text. Falls back to the .pptx when there is no PDF or pypdfium2 is
    missing.
    """
    if pdf_path is not None and Path(pdf_path).is_file():
        try:
            return pdf_slide_texts(pdf_path)
        except ImportError:
            pass
    if pptx_path is None:
        raise ValueError("designed_texts needs a .pptx or a readable PDF")
    return pptx_slide_texts(pptx_path)


def _shape_texts(shape: Any) -> list[str]:
    """Text in a shape, including groups and tables, in reading order."""
    texts: list[str] = []
    if hasattr(shape, "shapes"):
        for child in shape.shapes:  # group shape
            texts += _shape_texts(child)
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        texts.append(shape.text_frame.text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            texts.append(" | ".join(cell.text for cell in row.cells))
    return texts


def pptx_slide_texts(pptx_path: str | Path) -> list[str]:
    """Visible text of every slide in a .pptx, one string per slide."""
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    return ["\n".join(t for shape in s.shapes for t in _shape_texts(shape))
            for s in prs.slides]


def pdf_slide_texts(pdf_path: str | Path) -> list[str]:
    """Text of every page in a PDF (needs pypdfium2), one string per page."""
    import pypdfium2 as pdfium  # type: ignore[import-not-found]

    pdf = pdfium.PdfDocument(str(pdf_path))
    try:
        texts: list[str] = []
        for i in range(len(pdf)):
            page = pdf[i]
            try:
                textpage = page.get_textpage()
                texts.append(textpage.get_text_range())
                textpage.close()
            finally:
                page.close()
        return texts
    finally:
        pdf.close()
