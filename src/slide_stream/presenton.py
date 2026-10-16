"""Presenton integration: designed decks from an external Presenton server.

Presenton (https://github.com/presenton/presenton, Apache-2.0) is an
open-source AI presentation generator with designed templates. SlideStream
does not bundle or fork it: the user runs their own Presenton (Docker or
desktop) and this client calls its generation API
(``POST /api/v1/ppt/presentation/generate`` with one ``slides_markdown``
entry per SlideStream slide), then downloads the exported deck.

The deck is requested with no added title or table-of-contents slides so it
stays 1:1 with the source slides — narration and speaker notes stay aligned
by position. Presenton may rephrase slide text to fit a layout; narration in
SlideStream is built from the source text, so the two can differ slightly.
To limit that, every request carries ``instructions`` asking Presenton to
stay faithful to the source (see ``DEFAULT_PRESENTON_INSTRUCTIONS``).
"""

import os
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from urllib.parse import urljoin

import requests

from .code_slides import (
    CODE_SLIDE_MODES,
    choose_code_slides,
    neighbour_frames,
    plain_theme,
    render_code_slide,
    rewrite_pptx_code_slides,
    sample_theme,
)
from .fidelity import (
    DEFAULT_MIN_KEPT,
    FIDELITY_MODES,
    FidelityReport,
    check_fidelity,
    designed_texts,
)

DEFAULT_PRESENTON_URL = "http://localhost:5001"
DEFAULT_PRESENTON_TEMPLATE = "general"
# Presentation generation runs several LLM calls server-side; be generous.
DEFAULT_TIMEOUT_SECONDS = 600.0
# Sent as Presenton's ``instructions`` field, which reaches both its layout
# choice and its per-slide writing. The narration is built from the source,
# so anything Presenton drops or invents makes voice and picture disagree.
# The live trial (docs/PRESENTON_ANALYSIS.md) saw code replaced by prose, an
# invented person on a title slide and an invented table row.
DEFAULT_PRESENTON_INSTRUCTIONS = (
    "Stay faithful to each slide's Markdown; the slides will be narrated "
    "from that source text. Show code blocks exactly as written, line for "
    "line with their indentation; never describe or summarise code instead "
    "of showing it. Do not invent people, names, roles, quotes, statistics, "
    "dates or table rows that are not in the source; leave optional fields "
    "empty instead. Keep each slide's original title."
)


class PresentonError(Exception):
    """Raised when Presenton is unreachable or a generation request fails."""


@dataclass
class PresentonResult:
    """A generated deck downloaded from Presenton."""

    presentation_id: str | None
    path: Path
    edit_url: str | None = None


def resolve_presenton_settings(
    config: dict[str, Any],
    url: str | None = None,
    api_key: str | None = None,
    template: str | None = None,
    instructions: str | None = None,
) -> dict[str, Any]:
    """Resolve Presenton connection settings: CLI flag > config > env > default.

    ``instructions``: flag > ``providers.presenton.instructions`` > the
    built-in default. An empty string sends no instructions at all.

    Mirrors the layered story for other self-hosted servers (ADR-0002): the
    personal home config holds the URL/key once, flags override per run, and
    ``PRESENTON_API_KEY``/``PRESENTON_URL`` cover shell-only setups.
    """
    block = config.get("providers", {}).get("presenton", {}) or {}
    base_url = (
        url or block.get("base_url") or os.environ.get("PRESENTON_URL")
        or DEFAULT_PRESENTON_URL
    )
    return {
        "base_url": str(base_url).rstrip("/"),
        "api_key": (
            api_key or block.get("api_key")
            or os.environ.get("PRESENTON_API_KEY") or None
        ),
        "template": template or block.get("template") or DEFAULT_PRESENTON_TEMPLATE,
        "instructions": _resolve_instructions(instructions, block),
        # ADR-0013: rebuild slides whose code Presenton dropped, and report
        # how faithful the designed deck is to the source.
        "code_slides": _choice(block, "code_slides", CODE_SLIDE_MODES, "auto"),
        "fidelity": _choice(block, "fidelity", FIDELITY_MODES, "warn"),
        "fidelity_min_kept": _fraction(block, "fidelity_min_kept", DEFAULT_MIN_KEPT),
    }


def _choice(block: dict[str, Any], key: str, allowed: tuple[str, ...], default: str) -> str:
    value = str(block.get(key) or default).strip().lower()
    if value not in allowed:
        raise PresentonError(
            f"providers.presenton.{key} must be one of {', '.join(allowed)} "
            f"(got {block.get(key)!r})."
        )
    return value


def _fraction(block: dict[str, Any], key: str, default: float) -> float:
    raw = block.get(key, default)
    try:
        value = float(raw)
    except (TypeError, ValueError):
        raise PresentonError(
            f"providers.presenton.{key} must be a number from 0 to 1 (got {raw!r})."
        ) from None
    if not 0.0 <= value <= 1.0:
        raise PresentonError(
            f"providers.presenton.{key} must be from 0 to 1 (got {value})."
        )
    return value


def _resolve_instructions(flag: str | None, block: dict[str, Any]) -> str:
    """First of flag, config, default that is set; ``""`` means none."""
    for value in (flag, block.get("instructions")):
        if value is not None:
            return str(value).strip()
    return DEFAULT_PRESENTON_INSTRUCTIONS


def slides_to_markdown(slides: list[dict[str, Any]]) -> list[str]:
    """One ``slides_markdown`` entry per parsed slide (title + body).

    Markdown slides send their own Markdown body (``slide["markdown"]``), so
    code blocks, tables and sub-headings reach Presenton intact. Slides
    without one (PowerPoint input) send their content as bullets.

    A parser placeholder title (``Slide N``, ``has_real_title=False``) is not
    sent as a heading; a fully empty slide still gets an entry so the count
    never shifts against narration.
    """
    entries: list[str] = []
    for i, slide in enumerate(slides):
        title = str(slide.get("title", "")).strip()
        if not slide.get("has_real_title", True):
            title = ""
        body = slide.get("markdown")
        if isinstance(body, str):
            body_lines = [body.strip("\n")] if body.strip() else []
        else:
            body_lines = [
                f"- {str(c).strip()}"
                for c in slide.get("content", []) if str(c).strip()
            ]
        lines: list[str] = []
        if title:
            lines.append(f"# {title}")
        if body_lines:
            if lines:
                lines.append("")
            lines.extend(body_lines)
        entries.append("\n".join(lines) if lines else f"# Slide {i + 1}")
    return entries


class PresentonClient:
    """Minimal client for a self-hosted Presenton generation API."""

    def __init__(
        self,
        base_url: str,
        api_key: str | None = None,
        timeout: float = DEFAULT_TIMEOUT_SECONDS,
        session: Any = None,
        instructions: str = DEFAULT_PRESENTON_INSTRUCTIONS,
    ) -> None:
        self.base_url = base_url.rstrip("/")
        self.api_key = api_key
        self.instructions = instructions
        self.timeout = timeout
        self._session = session if session is not None else requests.Session()

    @property
    def _headers(self) -> dict[str, str]:
        if self.api_key:
            return {"Authorization": f"Bearer {self.api_key}"}
        return {}

    def _post(self, path: str, payload: dict[str, Any]) -> dict[str, Any]:
        url = f"{self.base_url}{path}"
        try:
            response = self._session.post(
                url, json=payload, headers=self._headers, timeout=self.timeout
            )
            response.raise_for_status()
        except requests.ConnectionError:
            raise PresentonError(
                f"could not reach Presenton at {self.base_url}. Start it "
                "(https://docs.presenton.ai) or point providers.presenton."
                "base_url / --presenton-url at your server."
            ) from None
        except requests.HTTPError as e:
            detail = ""
            if e.response is not None:
                try:
                    detail = str(e.response.json().get("detail", ""))
                except Exception:
                    detail = (e.response.text or "").strip()
            suffix = f": {detail}" if detail else ""
            raise PresentonError(
                f"Presenton request to {path} failed ({e}){suffix}"
            ) from None
        except requests.RequestException as e:
            raise PresentonError(f"Presenton request to {path} failed: {e}") from None
        data = response.json()
        if not isinstance(data, dict):
            raise PresentonError(f"Presenton returned an unexpected response: {data!r}")
        return data

    def _download(self, result: dict[str, Any], dest: Path) -> PresentonResult:
        raw_path = str(result.get("path", "")).strip()
        if not raw_path:
            raise PresentonError(
                "Presenton response did not include an exported file path."
            )
        # ``path`` is either an absolute URL or a server-relative path such as
        # /app_data/<id>/Deck.pptx, served from the Presenton origin.
        url = (
            raw_path if raw_path.startswith(("http://", "https://"))
            else urljoin(f"{self.base_url}/", raw_path.lstrip("/"))
        )
        try:
            response = self._session.get(
                url, headers=self._headers, timeout=self.timeout
            )
            response.raise_for_status()
        except requests.RequestException as e:
            raise PresentonError(f"could not download the deck from {url}: {e}") from None
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(response.content)
        edit_path = str(result.get("edit_path", "")).strip()
        return PresentonResult(
            presentation_id=(
                str(result["presentation_id"]) if result.get("presentation_id") else None
            ),
            path=dest,
            edit_url=(
                edit_path if edit_path.startswith(("http://", "https://"))
                else urljoin(f"{self.base_url}/", edit_path.lstrip("/"))
                if edit_path else None
            ),
        )

    def generate(
        self,
        slides_markdown: list[str],
        template: str,
        export_as: str,
    ) -> dict[str, Any]:
        """POST a generation request; returns Presenton's response JSON."""
        payload: dict[str, Any] = {
            "content": "",
            "slides_markdown": slides_markdown,
            "template": template,
            "export_as": export_as,
            # Keep the deck 1:1 with the source: no extra slides.
            "include_title_slide": False,
            "include_table_of_contents": False,
        }
        if self.instructions:
            payload["instructions"] = self.instructions
        return self._post("/api/v1/ppt/presentation/generate", payload)

    def export(self, presentation_id: str, export_as: str) -> dict[str, Any]:
        """Re-export an existing presentation in another format."""
        return self._post(
            f"/api/v1/ppt/presentation/{presentation_id}/export",
            {"export_as": export_as},
        )

    def generate_deck(
        self,
        slides: list[dict[str, Any]],
        template: str,
        export_as: str,
        dest: Path,
    ) -> PresentonResult:
        """Generate one deck (``pptx`` or ``pdf``) and download it to ``dest``."""
        result = self.generate(slides_to_markdown(slides), template, export_as)
        return self._download(result, Path(dest))

    def export_deck(self, presentation_id: str, export_as: str, dest: Path) -> PresentonResult:
        """Download an existing presentation in another format."""
        return self._download(self.export(presentation_id, export_as), Path(dest))

    def generate_deck_pair(
        self,
        slides: list[dict[str, Any]],
        template: str,
        pptx_dest: Path,
        pdf_dest: Path,
    ) -> PresentonResult:
        """Generate once, download both the editable PPTX and a PDF of it.

        The video pipeline rasterises the PDF; the user keeps the PPTX. One
        generation (plus a cheap re-export) keeps the two identical.
        """
        result = self.generate_deck(slides, template, "pptx", pptx_dest)
        if not result.presentation_id:
            raise PresentonError(
                "Presenton response did not include a presentation_id, so the "
                "PDF for the video cannot be exported."
            )
        pdf_result = self.export(result.presentation_id, "pdf")
        self._download(pdf_result, Path(pdf_dest))
        return result


def inject_speaker_notes(pptx_path: str | Path, notes: list[str]) -> None:
    """Write presenter notes into a deck, in slide order, in place.

    Blank entries leave a slide's existing notes untouched; notes beyond the
    deck's slide count are ignored.
    """
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    for slide, note in zip(prs.slides, notes, strict=False):
        text = str(note).strip()
        if not text:
            continue
        frame = slide.notes_slide.notes_text_frame
        if frame is not None:
            frame.text = text
    prs.save(str(pptx_path))


@dataclass
class DesignReview:
    """What ``review_designed_deck`` found and changed."""

    report: FidelityReport
    rebuilt: list[int] = field(default_factory=list)  # 0-based code slides rebuilt
    failures: list[str] = field(default_factory=list)  # reasons strict mode fails

    def to_dict(self) -> dict[str, Any]:
        return {
            **self.report.to_dict(),
            "code_slides_rebuilt": [i + 1 for i in self.rebuilt],
            "strict_failures": self.failures,
        }

    def write_json(self, path: str | Path) -> Path:
        import json

        dest = Path(path)
        dest.write_text(json.dumps(self.to_dict(), indent=2), encoding="utf-8")
        return dest


def review_designed_deck(
    slides: list[dict[str, Any]],
    settings: dict[str, Any],
    config: dict[str, Any],
    *,
    pptx_path: Path | None,
    pdf_path: Path | None,
    frames: list[Path] | None = None,
    frame_size: tuple[int, int] | None = None,
    work_dir: Path | None = None,
) -> DesignReview:
    """Check a designed deck against its source and rebuild lost code slides
    (ADR-0013).

    - Compares each slide's designed text (PDF preferred) with the source.
    - Picks slides to rebuild per ``settings["code_slides"]``.
    - Rewrites those slides in ``pptx_path`` (if given), and, when ``frames``
      are given, replaces their video frames in place with code slides drawn
      into ``work_dir``.
    - Colours come from the neighbouring frames, else SlideStream's own look.
    - Records strict-mode failures: code still missing, or too little text.
    """
    report = check_fidelity(
        slides, designed_texts(pptx_path, pdf_path),
        min_kept=float(settings.get("fidelity_min_kept", DEFAULT_MIN_KEPT)),
    )
    can_edit = pptx_path is not None or frames is not None
    chosen = (
        choose_code_slides(slides, report.code_missing(), settings.get("code_slides", "auto"))
        if can_edit else []
    )
    review = DesignReview(report=report, rebuilt=chosen)
    replaced = set(chosen)
    for index in chosen:
        theme = (
            sample_theme(neighbour_frames(index, frames, replaced)) if frames else None
        ) or plain_theme(config)
        if pptx_path is not None:
            rewrite_pptx_code_slides(pptx_path, slides, [index], theme)
        if frames is not None and frame_size is not None and work_dir is not None:
            work_dir.mkdir(parents=True, exist_ok=True)
            frames[index] = render_code_slide(
                work_dir / f"code_slide_{index + 1}.png", slides[index], theme, frame_size
            )
    still_missing = [i for i in report.code_missing() if i not in replaced]
    if still_missing:
        review.failures.append(
            "code missing on slide(s) " + ", ".join(str(i + 1) for i in still_missing)
        )
    low = report.below_min_kept()
    if low:
        review.failures.append(
            f"less than {report.min_kept:.0%} of the text kept on slide(s) "
            + ", ".join(str(i + 1) for i in low)
        )
    return review
